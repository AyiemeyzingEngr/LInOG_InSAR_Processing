from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image
import copy
SRC='/tmp/claude-0/fc/fc.pptx'
OUT='/home/user/LInOG_InSAR_Processing/linog_data/04_methods_notes/FBS FBD/LInOG_v4.0_Pipeline_Flowcharts_20261002_draft.pptx'
A='/tmp/claude-0/pp/'
prs=Presentation(SRC); blank=[l for l in prs.slide_layouts if l.name=='Blank'][0]
NAVY=RGBColor(0x1b,0x4f,0x72); NEWE=RGBColor(0xca,0x6f,0x1e); NEWF=RGBColor(0xfd,0xf2,0xe9); STDF=RGBColor(0xee,0xf5,0xfb)
OUTF=RGBColor(0xf4,0xfb,0xf6); OUTE=RGBColor(0x1e,0x84,0x49); GATE=RGBColor(0x11,0x7a,0x65); GREY=RGBColor(0x55,0x55,0x55); DARK=RGBColor(0x22,0x22,0x22); WHITE=RGBColor(255,255,255)
def box(sl,x,y,w,h,fill,line,lw=1.25,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.08):
    s=sl.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill; s.line.color.rgb=line; s.line.width=Pt(lw); s.shadow.inherit=False
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0]=radius
    s.text_frame.text=''; return s
def text(sl,x,y,w,h,parts,size=9,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,wrap=True):
    tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=wrap
    tf.margin_left=tf.margin_right=Inches(0.03); tf.margin_top=tf.margin_bottom=Inches(0.01); tf.vertical_anchor=anchor
    first=True
    for para in parts:  # para: list of (text, dict)
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False; p.alignment=align
        for t,st in para:
            r=p.add_run(); r.text=t; f=r.font; f.size=Pt(st.get('size',size)); f.bold=st.get('b',False); f.italic=st.get('i',False)
            f.color.rgb=st.get('color',DARK); f.name=st.get('font','Calibri')
    return tb
def tag(sl,x,y,t,fill,size=7,w=None):
    w=w or (0.12+0.075*len(t)*size/7)
    s=box(sl,x,y,w,0.19,fill,fill,radius=0.25); tf=s.text_frame; tf.margin_left=tf.margin_right=Inches(0.02); tf.margin_top=tf.margin_bottom=0
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; r=p.add_run(); r.text=t; r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name='Calibri'
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE; return s
def arrow(sl,x1,y1,x2,y2,color=NAVY,w=1.75):
    c=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w)
    ln=c.line._get_or_add_ln(); te=ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}); ln.append(te)
    return c
def pic(sl,path,x,y,w,h):
    iw,ih=Image.open(path).size; s=min(w/iw,h/ih); dw,dh=iw*s,ih*s
    return sl.shapes.add_picture(path,Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),Inches(dw),Inches(dh))
def header(sl,title,sub):
    text(sl,0.4,0.18,12.5,0.45,[[(title,{'b':True,'size':22,'color':NAVY})]])
    text(sl,0.4,0.66,12.5,0.3,[[(sub,{'size':11,'color':GREY})]])
PH=[('0','Set-up and pre-flight','Install ISCE2, MintPy, GDAL, SNAPHU; check provenance; size the machine; count the archive','gear',False,'provenance check',
     'Verified processing environment','Software checked, parallel settings sized, scene count matched','TERM','Summary of the pre-flight record: all 14 checks pass'),
    ('1','Unpack and unify FBS + FBD','Unzip L1.1 granules; beam check; sort FBS / FBD; band-limit FBS to 14 MHz at native 4.68 m range sampling','merge',True,'gates G1–G6',
     'One unified SLC stack','All dates at one 14 MHz bandwidth; 19–24 dates per frame','thumb_1.png','Range spectra: FBS before / after band-limiting match FBD'),
    ('2','DEM and network design','Shared DEM; pairs with temporal baseline ≤ 730 d and perpendicular baseline ≤ 1,500 m','mountain',False,'gate G17',
     'DEM + interferogram network','Pair list and run files (123–180 pairs formed per frame)','thumb_2.png','Baseline network of the pairs, 2007–2011'),
    ('3','Coregistration','Align every SLC to one FBS reference date (4.68 m range grid; FBD resampled onto it)','layers',False,'gates G7–G9, G16, G21',
     'Coregistered SLC stack','All acquisitions on the reference geometry','thumb_3.png','Coregistration residuals per date'),
    ('4','Interferogram generation','Form interferograms; multilook 28 × 12 (~90 m × 90 m); filter; unwrap with SNAPHU','fringe',False,'gates G10, G11, G18, G20',
     'Unwrapped interferograms','Interferograms and coherence for every admitted pair','thumb_4.png','Filtered interferogram, 8 Jan – 23 Feb 2008'),
    ('4.5','Visual quality control','Phase and amplitude report pages for every interferogram','grid',False,'visual check',
     'Interferogram report grids','Poor pairs flagged and excluded','thumb_5.png','Report page used to inspect every pair'),
    ('5','SBAS time-series analysis','MintPy inversion; DEM-error correction; ramp removal; coherence and water masks','chart',False,'gates G12, G13, G15, G19, G22, G23',
     'Displacement time series','LOS velocity, displacement per date, temporal coherence','thumb_6.png','Network coloured by coherence; coherence matrix'),
    ('6','Geocoding and delivery','Geocode to a common 1/1200° (~90 m) grid; 16 products per frame; 27-check verification','map',False,'gate G14; 27 checks',
     '16 geocoded products per frame','LOS, vertical, horizontal GeoTIFFs; KMZ; hillshade maps','thumb_7.png','Delivered LOS velocity map (cm/yr)'),
    ('7','Interpretation','Read the velocity maps; measure subsidence against nearby stable ground','lens',False,'two-track check',
     'Ground-deformation results','e.g. Cabanatuan 4.4–4.7 cm/yr vertical, two tracks','thumb_8.png','Cabanatuan bowl, Path 448 and Path 449')]
def grid_slide(real):
    sl=prs.slides.add_slide(blank)
    header(sl,'LInOG FBS + FBD InSAR processing workflow'+(': real outputs' if real else ''),
           'ALOS-1 PALSAR, pipeline v4.0, one frame: each phase (left) and its output (right); read left to right, row by row.'+(' Sample: Path 449 Frame 0280.' if real else ''))
    CW,CH=4.05,1.86; X0,Y0,GX,GY=0.35,1.05,0.22,0.17
    pos=[]
    for i,ph in enumerate(PH):
        r,cidx=divmod(i,3); x=X0+cidx*(CW+GX); y=Y0+r*(CH+GY); pos.append((x,y))
        num,title,body,icon,new,gate,ot,ob,img,cap=ph
        ec=NEWE if new else NAVY
        PW=2.35 if not real else 2.2
        box(sl,x,y,PW,CH,NEWF if new else STDF,ec,2.0 if new else 1.25)
        pic(sl,A+'icon_%s.png'%icon,x+0.07,y+0.08,0.55,0.55)
        text(sl,x+0.66,y+0.06,PW-0.7,0.22,[[('Phase '+num,{'b':True,'size':9,'color':ec})]])
        text(sl,x+0.66,y+0.27,PW-0.7,0.42,[[(title,{'b':True,'size':10.5,'color':NAVY})]])
        text(sl,x+0.08,y+0.72,PW-0.14,0.95,[[(body,{'size':8})]])
        text(sl,x+0.08,y+CH-0.27,PW-0.14,0.22,[[(gate,{'i':True,'size':7.5,'color':GATE})]])
        if new: tag(sl,x+PW-0.92,y+CH-0.47,'NEW IN YEAR 2',NEWE,size=6.5,w=0.86)
        ox=x+PW+0.17; OW=CW-PW-0.17
        box(sl,ox,y,OW,CH,OUTF,OUTE,1.25)
        arrow(sl,x+PW+0.01,y+CH/2,ox-0.01,y+CH/2,OUTE,1.5)
        tag(sl,ox+0.07,y+0.07,'OUTPUT',OUTE,size=6.5,w=0.55)
        if not real:
            text(sl,ox+0.06,y+0.32,OW-0.1,0.55,[[(ot,{'b':True,'size':9,'color':OUTE})]])
            text(sl,ox+0.06,y+0.9,OW-0.1,0.95,[[(ob,{'size':8})]])
        else:
            if img=='TERM':
                t=box(sl,ox+0.06,y+0.3,OW-0.12,1.12,RGBColor(0x1e,0x1e,0x1e),RGBColor(0x44,0x44,0x44),0.75,shape=MSO_SHAPE.RECTANGLE)
                lines=['# pre-flight record (felix, 17 Sep 2026)','provenance ... OK','ISCE2 · MintPy · GDAL · SNAPHU OK','48 cores, 125 GB RAM','779 ALOS-1 L1.1 zips','RESULT: PASS 14/14']
                text(sl,ox+0.08,y+0.32,OW-0.16,1.08,[[(l,{'size':6,'font':'Consolas','color':(RGBColor(0x7d,0xce,0xa0) if 'PASS' in l else RGBColor(0xea,0xea,0xea))})] for l in lines])
            else:
                pic(sl,A+img,ox+0.06,y+0.3,OW-0.12,1.12)
            text(sl,ox+0.04,y+1.45,OW-0.08,0.48,[[(cap,{'size':7})]])
    for i in range(len(PH)-1):
        (x1,y1),(x2,y2)=pos[i],pos[i+1]
        if y1==y2: arrow(sl,x1+CW+0.01,y1+CH/2,x2-0.01,y2+CH/2,NAVY,2)
        else:
            gy=y1+CH+GY/2; xs=x1+CW/2; xe=x2+CW/2
            fb=sl.shapes.build_freeform(Inches(xs),Inches(y1+CH),scale=1.0)
            fb.add_line_segments([(Inches(xs),Inches(gy)),(Inches(xe),Inches(gy)),(Inches(xe),Inches(y2))],close=False)
            ff=fb.convert_to_shape(); ff.fill.background(); ff.line.color.rgb=NAVY; ff.line.width=Pt(1.5)
            ln=ff.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
    # legend
    ly=7.1
    for j,(f,e,l) in enumerate([(NEWF,NEWE,'new for FBS+FBD processing'),(STDF,NAVY,'carried over from FBS-only'),(OUTF,OUTE,'output of the phase')]):
        box(sl,0.4+j*2.6,ly,0.3,0.16,f,e,1.25); text(sl,0.75+j*2.6,ly-0.03,2.2,0.22,[[(l,{'size':8,'color':GREY})]])
    text(sl,8.3,ly-0.03,4.8,0.22,[[('Green italic: automated quality gates (G1–G23); a failed gate stops the phase.',{'i':True,'size':8,'color':GATE})]])
    return sl
grid_slide(False); grid_slide(True)
# ---- mosaic progression slide ----
sl=prs.slides.add_slide(blank)
header(sl,'Progression of LInOG ground-velocity coverage, 2007–2011','ALOS-1 PALSAR ascending tracks, deramped LOS velocity, same ±5 cm/yr colour scale')
for k,(x,t1,t1c,t2,t3,img) in enumerate([(0.55,'YEAR 1 (reported June 2026)',RGBColor(0x7f,0x8c,0x8d),'FBS only · Path 448 · 5 frames','10 dates/frame · 34–36 interferograms/frame\n~10,000 km² reported · no water mask','panel_y1.png'),
                                         (6.6,'YEAR 2, QUARTER 2 (this report)',NEWE,'FBS + FBD · Paths 447–449 · 12 frames','19–24 dates/frame · 1,634 interferograms in total\n27,893 km² with valid data · water-masked (v1.2)','panel_q2.png')]):
    s=box(sl,x+0.9,1.05,3.6,0.36,t1c,t1c,radius=0.3); tf=s.text_frame; p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; r=p.add_run(); r.text=t1; r.font.size=Pt(13); r.font.bold=True; r.font.color.rgb=WHITE; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    text(sl,x,1.45,5.4,0.3,[[(t2,{'b':True,'size':12,'color':NAVY})]],align=PP_ALIGN.CENTER)
    pic(sl,A+img,x,1.8,5.4,4.75)
    st=box(sl,x+0.25,6.62,4.9,0.62,WHITE,RGBColor(0xbb,0xbb,0xbb),1)
    text(sl,x+0.3,6.65,4.8,0.58,[[(l,{'size':10})] for l in t3.split('\n')])
a=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(5.95),Inches(3.75),Inches(0.6),Inches(0.45)); a.fill.solid(); a.fill.fore_color.rgb=NEWE; a.line.color.rgb=NEWE
pic(sl,A+'colorbar.png',12.05,2.3,1.2,3.6)
# ---- GeoLab tutorial route slide ----
sl=prs.slides.add_slide(blank)
header(sl,'LInOG tutorial route on EarthScope GeoLab (manual v4.0, Chapter 11)','Browser-based JupyterLab server: trainees run the LInOG notebook on a four-date demonstration stack, without access to the project server')
GS=[('1','EarthScope account','Google log-in to the GeoLab hub','gear'),('2','Start a server','29 GB profile: 3.66 CPU cores, 29,776 MiB memory','layers'),
    ('3','One-time install','ISCE2, MintPy, GDAL, SNAPHU into home (~7 min)','archive'),('4','Demonstration stack','Path 449 Frame 0290: 4 dates, 6 pairs (2008)','mountain'),
    ('5','Run the notebook','Cells 0–18 call the LInOG scripts, Phases 0–6','chart'),('6','Check and deliver','18-item run check; 16 products','map')]
bw,gap,y=1.92,0.24,1.1
for i,(n,t,b,ic) in enumerate(GS):
    x=0.45+i*(bw+gap); box(sl,x,y,bw,2.15,STDF,NAVY)
    text(sl,x+0.08,y+0.05,1.0,0.25,[[('Step '+n,{'b':True,'size':10,'color':NAVY})]])
    pic(sl,A+'icon_%s.png'%ic,x+bw/2-0.36,y+0.3,0.72,0.72)
    text(sl,x+0.05,y+1.08,bw-0.1,0.3,[[(t,{'b':True,'size':10,'color':NAVY})]],align=PP_ALIGN.CENTER)
    text(sl,x+0.08,y+1.4,bw-0.16,0.7,[[(b,{'size':9})]],align=PP_ALIGN.CENTER)
    if i<5: arrow(sl,x+bw+0.02,y+1.07,x+bw+gap-0.02,y+1.07)
text(sl,0.45,3.4,12.5,0.3,[[('REAL OUTPUTS PRODUCED ON GEOLAB (September 2026; four-date demonstration stack, for training only)',{'b':True,'size':12,'color':OUTE})]],align=PP_ALIGN.CENTER)
ow=4.05
for k,(f,c) in enumerate([('s35_1.jpg','Phase 4.5: interferogram report sheet (GeoLab, 20 Sept 2026)'),('s35_2.png','Phase 5: velocity in radar coordinates (GeoLab, 22 Sept 2026)'),('s35_3.jpg','Phase 6: delivered velocity map, cm/yr (GeoLab, 20 Sept 2026)')]):
    x=0.45+k*(ow+0.2); box(sl,x,3.8,ow,3.45,OUTF,OUTE)
    pic(sl,A+f,x+0.12,3.9,ow-0.24,2.85)
    text(sl,x+0.1,6.8,ow-0.2,0.4,[[(c,{'size':9})]],align=PP_ALIGN.CENTER)
prs.save(OUT); print('saved',OUT,len(prs.slides))
